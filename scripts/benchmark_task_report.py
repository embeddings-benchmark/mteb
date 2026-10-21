"""Combine the benchmark task analyses into one xlsx report with flags, descriptions and sample examples.

Reads the CSVs written by the other `benchmark_*` scripts from `--analysis-dir`:
- `spread.csv` (`benchmark_score_spread.py`)
- `concentration.csv` and `concentration_0.05.csv` (`benchmark_score_concentration.py` with `--tolerance 0.02` and `0.05`)
- `noise.csv` (`benchmark_task_noise.py`)
- `corr/task_correlation.csv` and `corr/task_groups.csv` (`benchmark_task_correlation.py`)

Every task gets flags and is listed on the `Tasks` sheet (sorted by the number of flags) together with its
description and the first rows of the dataset at the pinned revision, loaded with `datasets` in streaming mode.

Sheets:
- `Tasks`: every task in the benchmark with flags, metrics, description and examples.
- `Examples`: the fetched dataset rows, one per line.
- `Correlation groups`: groups of highly correlated tasks.
- `Flags`: definitions of the flags.

Example:
    uv run --no-sync python scripts/benchmark_task_report.py --analysis-dir analysis --output benchmark_task_report.xlsx
"""

from __future__ import annotations

import argparse
import itertools
import json
from pathlib import Path
from typing import Any

import datasets
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import mteb

RETRIEVAL_TYPES = {
    "Retrieval",
    "Reranking",
    "InstructionRetrieval",
    "InstructionReranking",
}


def _pick_configs(configs: list[str], task_type: str) -> list[str]:
    """Pick the configs to show: queries and corpus for retrieval, one (preferably English) subset otherwise."""

    def prefer_english(names: list[str]) -> list[str]:
        english = [n for n in names if "eng" in n or n.startswith("en") or "-en" in n]
        default = [n for n in names if n == "default"]
        return (default or english or names)[:1]

    if task_type in RETRIEVAL_TYPES:
        picked = []
        for kind in ("queries", "corpus"):
            names = [c for c in configs if kind in c and "instruction" not in c]
            picked += prefer_english(names) if names else []
        if picked:
            return picked
        configs = [c for c in configs if "qrels" not in c and "top_ranked" not in c]
    return prefer_english(configs) if configs else []


def _format_row(row: dict[str, Any], max_chars: int) -> str:
    parts = []
    for key, value in row.items():
        if key in ("id", "_id", "query-id", "corpus-id", "idx") or value is None:
            continue
        text = (
            value
            if isinstance(value, str)
            else json.dumps(value, ensure_ascii=False, default=str)
        )
        text = " ".join(text.split())
        if len(text) > max_chars:
            text = text[:max_chars] + "…"
        parts.append(f"{key}: {text}")
    return " | ".join(parts)


def fetch_examples(
    task: mteb.AbsTask, n_rows: int, max_chars: int
) -> list[dict[str, str]]:
    path = task.metadata.dataset["path"]
    revision = task.metadata.dataset.get("revision")
    eval_split = task.metadata.eval_splits[0]
    try:
        configs = datasets.get_dataset_config_names(path, revision=revision)
    except Exception as e:
        print(f"{task.metadata.name}: failed to list configs: {e}")
        return []
    examples = []
    for config in _pick_configs(configs, task.metadata.type):
        try:
            # older retrieval datasets name the split after the config (e.g. `corpus/corpus`)
            split_names = datasets.get_dataset_split_names(
                path, config, revision=revision
            )
            split = eval_split if eval_split in split_names else split_names[0]
            dataset = datasets.load_dataset(
                path, config, split=split, revision=revision, streaming=True
            )
            rows = list(itertools.islice(dataset, n_rows))
        except Exception as e:
            print(f"{task.metadata.name}: failed to load {config}: {e}")
            continue
        for row in rows:
            examples.append(
                {
                    "task": task.metadata.name,
                    "dataset": path,
                    "config": config,
                    "split": split,
                    "row": _format_row(row, max_chars),
                }
            )
    return examples


def build_metrics(analysis_dir: Path, n_least_correlated: int) -> pd.DataFrame:
    spread = pd.read_csv(analysis_dir / "spread.csv", index_col="task")
    conc = pd.read_csv(analysis_dir / "concentration.csv", index_col="task")
    conc05 = pd.read_csv(analysis_dir / "concentration_0.05.csv", index_col="task")
    noise = pd.read_csv(analysis_dir / "noise.csv", index_col="task")
    corr = pd.read_csv(analysis_dir / "corr" / "task_correlation.csv", index_col=0)

    off_diagonal = corr.where(
        ~pd.DataFrame(
            [[a == b for b in corr.columns] for a in corr.index],
            index=corr.index,
            columns=corr.columns,
        )
    )
    correlation = pd.DataFrame(
        {
            "mean_corr": off_diagonal.mean(),
            "max_corr": off_diagonal.max(),
            "most_similar": off_diagonal.idxmax(),
        }
    )

    metrics = pd.concat(
        [
            noise[["type", "main_score", "n_items", "n_subsets"]],
            spread[
                [
                    "top1_model",
                    "top1",
                    "top10",
                    "top50",
                    "diff_1_50",
                    "jump_1_10",
                    "random",
                ]
            ].rename(columns={"diff_1_50": "top1-top50", "jump_1_10": "top1-top10"}),
            conc[["median", "within", "n_breakout", "breakout_models"]].rename(
                columns={"within": "within_0.02"}
            ),
            conc05[["within"]].rename(columns={"within": "within_0.05"}),
            noise[
                [
                    "se_diff",
                    "share_distinguishable",
                    "share_distinguishable_conservative",
                ]
            ],
            correlation,
        ],
        axis=1,
    )

    least_correlated = set(metrics["mean_corr"].nsmallest(n_least_correlated).index)
    flags = {
        "concentrated_0.02": lambda t, r: r["within_0.02"] >= 0.5,
        "concentrated_0.05": lambda t, r: r["within_0.05"] >= 0.8,
        "low_diversity": lambda t, r: r["top1-top50"] <= 0.05,
        "top_identical": lambda t, r: r["top1-top50"] <= 0.02,
        "outlier_jump": lambda t, r: r["top1-top10"] >= 0.25,
        "least_correlated": lambda t, r: t in least_correlated,
        "noisy": lambda t, r: r["share_distinguishable"] < 0.5,
    }
    metrics["flags"] = [
        ", ".join(name for name, check in flags.items() if check(task, row))
        for task, row in metrics.iterrows()
    ]
    metrics["n_flags"] = metrics["flags"].str.count(",") + (metrics["flags"] != "")
    return metrics


FLAG_DEFINITIONS = [
    (
        "concentrated_0.02",
        "within_0.02 >= 0.5",
        "Share of models (all models covering at least 97% of the benchmark) scoring within ±0.02 of the median.",
        "Most models get practically the same score, so the task barely affects the ranking. Usually the task is "
        "at chance for everyone or saturated; a few breakout models often indicate training on the task data.",
    ),
    (
        "concentrated_0.05",
        "within_0.05 >= 0.8",
        "Share of models scoring within ±0.05 of the median.",
        "Same as concentrated_0.02 with a wider band.",
    ),
    (
        "low_diversity",
        "top1-top50 <= 0.05",
        "Score of the best model minus the score of the 50th-best model.",
        "The top 50 models score almost the same, so the task does not separate strong models (saturated or at chance).",
    ),
    (
        "top_identical",
        "top1-top50 <= 0.02",
        "Score of the best model minus the score of the 50th-best model.",
        "Stricter low_diversity: the best and the 50th-best model are practically identical (e.g. BUCC.v2, 0.005).",
    ),
    (
        "outlier_jump",
        "top1-top10 >= 0.25",
        "Score of the best model minus the score of the 10th-best model.",
        "One or a few models score far above the rest, which often means they were trained on the task data.",
    ),
    (
        "least_correlated",
        "mean_corr among the 20 lowest",
        "Mean Spearman correlation of the task's ranking of the top 50 models with every other task.",
        "The task ranks models unlike any other task: either noise or a unique signal, check the other flags.",
    ),
    (
        "noisy",
        "share_distinguishable < 0.5",
        "Share of top-50 model pairs whose score difference exceeds 1.96 * se_diff, where se_diff is the estimated "
        "standard error of a difference from test-set size (and evaluation randomness for classification and clustering).",
        "Most differences between strong models are within sampling noise, usually because the test set is small.",
    ),
]


def write_sheet(
    writer: pd.ExcelWriter, df: pd.DataFrame, name: str, wide: dict[str, int]
) -> None:
    df.to_excel(writer, sheet_name=name, index=False)
    sheet = writer.sheets[name]
    header_fill = PatternFill("solid", start_color="DDEBF7")
    for cell in sheet[1]:
        cell.font = Font(name="Arial", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial")
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if isinstance(cell.value, float):
                cell.number_format = "0.000"
    for i, column in enumerate(df.columns, start=1):
        sheet.column_dimensions[get_column_letter(i)].width = wide.get(
            column, max(10, min(28, len(str(column)) + 2))
        )
    sheet.freeze_panes = "B2"
    sheet.auto_filter.ref = sheet.dimensions


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--benchmark", default="MTEB(Multilingual, v2)")
    parser.add_argument("--analysis-dir", type=Path, required=True)
    parser.add_argument(
        "--output", type=Path, default=Path("benchmark_task_report.xlsx")
    )
    parser.add_argument("--n-least-correlated", type=int, default=20)
    parser.add_argument("--n-rows", type=int, default=3)
    parser.add_argument("--max-chars", type=int, default=200)
    args = parser.parse_args()

    metrics = build_metrics(args.analysis_dir, args.n_least_correlated)
    tasks = {t.metadata.name: t for t in mteb.get_benchmark(args.benchmark).tasks}
    metrics.insert(
        1, "description", [tasks[t].metadata.description for t in metrics.index]
    )
    metrics.insert(
        2,
        "eval_splits",
        [", ".join(tasks[t].metadata.eval_splits) for t in metrics.index],
    )
    metrics.insert(
        3,
        "n_languages",
        [len(tasks[t].metadata.languages) for t in metrics.index],
    )

    metrics = metrics.sort_values("n_flags", ascending=False)
    examples = []
    for name in metrics.index:
        task_examples = fetch_examples(tasks[name], args.n_rows, args.max_chars)
        print(f"{name}: {len(task_examples)} examples")
        examples.extend(task_examples)
    examples_df = pd.DataFrame(
        examples, columns=["task", "dataset", "config", "split", "row"]
    )
    metrics.insert(
        4,
        "examples",
        [
            "\n".join(examples_df.loc[examples_df["task"] == name, "row"])
            for name in metrics.index
        ],
    )

    groups = pd.read_csv(args.analysis_dir / "corr" / "task_groups.csv")
    groups["tasks"] = groups["tasks"].str.replace(", ", "\n")

    front = ["task", "flags", "n_flags", "type", "description", "examples"]
    all_tasks = metrics.reset_index(names="task")
    all_tasks = all_tasks[front + [c for c in all_tasks.columns if c not in front]]

    wide = {
        "task": 34,
        "flags": 30,
        "description": 50,
        "examples": 80,
        "breakout_models": 50,
        "top1_model": 30,
        "most_similar": 30,
        "row": 120,
        "tasks": 45,
        "condition": 30,
        "metric": 70,
        "meaning": 80,
    }
    flag_definitions = pd.DataFrame(
        FLAG_DEFINITIONS, columns=["flag", "condition", "metric", "meaning"]
    )
    flag_definitions["n_tasks"] = [
        all_tasks["flags"].str.split(", ").apply(lambda f, flag=flag: flag in f).sum()
        for flag in flag_definitions["flag"]
    ]

    args.output.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(args.output, engine="openpyxl") as writer:
        write_sheet(writer, all_tasks, "Tasks", wide)
        write_sheet(writer, examples_df, "Examples", wide)
        write_sheet(writer, groups, "Correlation groups", wide)
        write_sheet(
            writer,
            flag_definitions,
            "Flags",
            wide,
        )
    n_flagged = int((all_tasks["n_flags"] > 0).sum())
    print(f"{len(all_tasks)} tasks ({n_flagged} flagged), wrote {args.output}")


if __name__ == "__main__":
    main()
